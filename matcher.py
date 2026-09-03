"""
LPN Matcher - Chedraui (logica principal)
==========================================
Empareja el "Codigo" del Excel de pedido con el LPN del PDF de etiquetas.
Extraido de lpn_matcher.py para ser usado como modulo dentro de la app web.
"""

import io
import os
import re
import subprocess
import unicodedata
from collections import Counter

import pymupdf as fitz
import numpy as np
import openpyxl
import xlrd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image
from pyzbar.pyzbar import decode as decode_barcode
from rapidfuzz import fuzz
from scipy.optimize import linear_sum_assignment


def _png_bytes(pil_img):
    buf = io.BytesIO()
    pil_img.save(buf, format="PNG")
    return buf.getvalue()


def ocr_pdf(pdf_path, workdir, progress_cb=None):
    pages_dir = os.path.join(workdir, "pages")
    os.makedirs(pages_dir, exist_ok=True)
    # 200dpi es necesario: a 150dpi Tesseract confunde acentos (p.ej.
    # "Descripcion" -> "Descripcién"), lo que rompe la extraccion de campos
    # y empeora el emparejamiento. No bajar esta resolucion.
    subprocess.run(
        ["pdftoppm", "-png", "-r", "200", pdf_path, os.path.join(pages_dir, "pg")],
        check=True,
    )
    files = sorted(os.listdir(pages_dir))
    total = len(files)
    results = []
    for i, f in enumerate(files):
        if progress_cb:
            progress_cb(i, total)
        path = os.path.join(pages_dir, f)
        full_img = Image.open(path)

        # El LPN se lee decodificando el codigo de barras real (CODE128), no el
        # texto OCR impreso debajo: el OCR confunde digitos (6/8, 5/6, etc.)
        # con bastante frecuencia y el codigo de barras es exacto.
        lpn = None
        barcodes = decode_barcode(full_img)
        if barcodes:
            lpn = barcodes[0].data.decode()

        # El texto (Sku, Descripcion, UPC, ASN...) ocupa la mitad superior de
        # la hoja; recortar antes de OCRear reduce el tiempo de Tesseract sin
        # perder texto (0.45 deja margen de sobra: todo el bloque termina
        # antes del 45% de la altura de la pagina).
        crop = full_img.crop((0, 0, full_img.width, int(full_img.height * 0.45)))
        text = subprocess.run(
            ["tesseract", "-", "stdout", "--psm", "6", "-l", "spa"],
            input=_png_bytes(crop),
            capture_output=True,
        ).stdout.decode("utf-8", "ignore")

        if not lpn:
            # respaldo si el codigo de barras no se pudo leer (raro)
            lines = [l.strip() for l in text.splitlines() if l.strip()]
            if lines:
                m = re.match(r"^(\d{6,10})$", lines[0])
                if m:
                    lpn = m.group(1)

        def grab(label):
            mm = re.search(label + r"\s*:?\s*(.+)", text)
            return mm.group(1).strip() if mm else None

        dm = re.search(r"Descripcion:?\s*(.+?)\n(?:\s*(.+?)\n)?Proveedor", text, re.S)
        descripcion = None
        if dm:
            descripcion = re.sub(
                r"\s+", " ", (dm.group(1) or "") + " " + (dm.group(2) or "")
            ).strip()

        results.append(
            {
                "page": i + 1,
                "lpn_barcode": lpn,
                "orden_compra": grab(r"No\. Orden de Compra"),
                "cantidad_cajas": grab("Cantidad de Cajas"),
                "sku": grab("Sku"),
                "upc": grab("UPC"),
                "asn": grab("ASN"),
                "descripcion": descripcion,
            }
        )
    if progress_cb:
        progress_cb(total, total)
    return results


def _load_excel_xls(xls_path):
    wb = xlrd.open_workbook(xls_path)
    sh = wb.sheet_by_index(0)
    rows = []
    for r in range(1, sh.nrows):
        codigo = str(sh.cell_value(r, 0)).strip()
        if not codigo:
            continue
        rows.append(
            {
                "codigo": codigo,
                "ean13": str(sh.cell_value(r, 1)).strip(),
                "codigo_cliente": str(sh.cell_value(r, 2)).strip(),
                "descripcion": str(sh.cell_value(r, 3)).strip(),
                "unidad": str(sh.cell_value(r, 4)).strip(),
                "cantidad": sh.cell_value(r, 5),
                "cajas_pedido": str(sh.cell_value(r, 6)).strip(),
                "precio_lista": sh.cell_value(r, 7) if sh.ncols > 7 else None,
            }
        )
    return rows


def _load_excel_xlsx(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    sh = wb.worksheets[0]
    rows = []
    for row in sh.iter_rows(min_row=2, values_only=True):
        codigo = str(row[0]).strip() if row[0] is not None else ""
        if not codigo:
            continue
        rows.append(
            {
                "codigo": codigo,
                "ean13": str(row[1]).strip() if len(row) > 1 and row[1] is not None else "",
                "codigo_cliente": str(row[2]).strip() if len(row) > 2 and row[2] is not None else "",
                "descripcion": str(row[3]).strip() if len(row) > 3 and row[3] is not None else "",
                "unidad": str(row[4]).strip() if len(row) > 4 and row[4] is not None else "",
                "cantidad": row[5] if len(row) > 5 else None,
                "cajas_pedido": str(row[6]).strip() if len(row) > 6 and row[6] is not None else "",
                "precio_lista": row[7] if len(row) > 7 else None,
            }
        )
    return rows


def load_excel(path):
    """Lee el Excel de pedido (.xls o .xlsx), excluyendo filas sin Codigo."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        return _load_excel_xls(path)
    return _load_excel_xlsx(path)


def strip_upc(u):
    if not u:
        return None
    u = re.sub(r"\D", "", u)
    return u[1:] if len(u) == 14 else u


def norm(s):
    """Normaliza texto: sin acentos, mayusculas, sin relleno (empaque, colores surtidos...)."""
    s = s or ""
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = s.upper()
    s = re.sub(r"EMPAQUE\s*:?\s*\d+", "", s)
    s = re.sub(r"COLORES SURTID\w*", "", s)
    s = re.sub(r"MIX COLORES?", "", s)
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def size_token(s):
    """Detecta talla/tamano (Chico/S, Mediano/M, Grande/G) para reforzar el emparejamiento."""
    s = (s or "").upper()
    if "CHIC" in s:
        return "S"
    if "MEDIAN" in s:
        return "M"
    if "GRAND" in s:
        return "G"
    m = re.search(r"\bTALLA\s+([SMLG])\b", s)
    if m:
        return {"S": "S", "M": "M", "L": "G", "G": "G"}[m.group(1)]
    return None


def match(excel_rows, pdf_entries):
    """
    Empareja en 3 fases:
      1) SKU exacto (Sku del PDF == Codigo Cliente del Excel)
      2) UPC/EAN exacto
      3) Para lo que quede, asignacion OPTIMA GLOBAL (algoritmo hungaro) usando
         similitud de descripcion + bonus/penalizacion por talla/tamano.
    """
    n_excel, n_pdf = len(excel_rows), len(pdf_entries)
    used_pdf = set()
    results = [None] * n_excel

    for idx, e in enumerate(excel_rows):
        for j, p in enumerate(pdf_entries):
            if j in used_pdf:
                continue
            if p["sku"] == e["codigo_cliente"]:
                results[idx] = (p, "SKU exacto", 100)
                used_pdf.add(j)
                break

    for idx, e in enumerate(excel_rows):
        if results[idx]:
            continue
        for j, p in enumerate(pdf_entries):
            if j in used_pdf:
                continue
            if strip_upc(p["upc"]) == e["ean13"]:
                results[idx] = (p, "UPC/EAN", 100)
                used_pdf.add(j)
                break

    pending_idx = [idx for idx in range(n_excel) if results[idx] is None]
    pending_pdf = [j for j in range(n_pdf) if j not in used_pdf]

    if pending_idx and pending_pdf:
        n, m = len(pending_idx), len(pending_pdf)
        score = np.zeros((n, m))
        for a, idx in enumerate(pending_idx):
            e = excel_rows[idx]
            ne, se = norm(e["descripcion"]), size_token(e["descripcion"])
            for b, j in enumerate(pending_pdf):
                p = pdf_entries[j]
                npd, sp = norm(p["descripcion"] or ""), size_token(p["descripcion"] or "")
                base = fuzz.token_set_ratio(ne, npd)
                if se and sp:
                    base += 25 if se == sp else -40
                score[a, b] = base
        row_ind, col_ind = linear_sum_assignment(-score)
        for a, b in zip(row_ind, col_ind):
            idx, j = pending_idx[a], pending_pdf[b]
            s = score[a, b]
            p = pdf_entries[j]
            metodo = (
                "Descripcion (asignacion optima)"
                if s >= 40
                else "Descripcion (BAJA confianza - revisar)"
            )
            results[idx] = (p, metodo, round(float(s), 1))
            used_pdf.add(j)

    out = []
    for e, res in zip(excel_rows, results):
        if res:
            p, metodo, conf = res
            out.append(
                {
                    **e,
                    "LPN": p["lpn_barcode"],
                    "metodo": metodo,
                    "confianza": conf,
                    "pdf_sku": p["sku"],
                    "pdf_descripcion": p["descripcion"],
                    "pdf_pagina": p["page"],
                }
            )
        else:
            out.append(
                {
                    **e,
                    "LPN": None,
                    "metodo": "SIN COINCIDENCIA",
                    "confianza": 0,
                    "pdf_sku": None,
                    "pdf_descripcion": None,
                    "pdf_pagina": None,
                }
            )
    return out


def categoria_metodo(metodo):
    """Clasifica el metodo de match en una categoria de confianza para colorear."""
    if metodo == "SIN COINCIDENCIA":
        return "bad"
    if "BAJA" in metodo:
        return "warn"
    if metodo.startswith("Descripcion"):
        return "good"
    return "ok"


def write_output(results, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pedido + LPN"
    headers = [
        "Codigo", "Ean13", "Codigo Cliente", "Descripcion", "Unidad", "Cantidad",
        "Cajas Pedido", "Precio Lista", "LPN", "Metodo Match", "Confianza (%)",
        "PDF Sku (verif.)", "PDF Descripcion (verif.)", "PDF Pagina",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5496")
        cell.alignment = Alignment(horizontal="center")

    fills = {
        cat: PatternFill("solid", fgColor=c)
        for cat, c in zip(("ok", "good", "warn", "bad"), ("C6EFCE", "DDEBF7", "FFEB9C", "FFC7CE"))
    }
    for r in results:
        row = [
            r["codigo"], r["ean13"], r["codigo_cliente"], r["descripcion"], r["unidad"],
            r["cantidad"], r["cajas_pedido"], r["precio_lista"], r["LPN"], r["metodo"],
            r["confianza"], r["pdf_sku"], r["pdf_descripcion"], r["pdf_pagina"],
        ]
        ws.append(row)
        ridx = ws.max_row
        fill = fills[categoria_metodo(r["metodo"])]
        for c in range(1, len(row) + 1):
            ws.cell(row=ridx, column=c).fill = fill

    for i, w in enumerate([10, 15, 14, 42, 10, 10, 12, 12, 12, 20, 12, 14, 42, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Resumen")
    c = Counter(r["metodo"] for r in results)
    ws2.append(["Metodo", "Cantidad"])
    for k, v in c.items():
        ws2.append([k, v])
    ws2.append(["TOTAL", len(results)])
    for cell in ws2[1]:
        cell.font = Font(bold=True)

    wb.save(out_path)
    return dict(c)


def annotate_pdf(pdf_path, results, out_pdf_path):
    """
    Genera una copia del PDF de etiquetas con un recuadro en cada pagina
    mostrando el "Codigo" (columna A del Excel) que le corresponde, para
    poder imprimir la etiqueta y pegarla en la caja identificada por ese
    codigo (las cajas no estan identificadas por el Codigo Cliente).
    El color del recuadro sigue la misma confianza que el Excel de salida.
    """
    box_colors = {
        "ok": (0.78, 0.94, 0.81),
        "good": (0.87, 0.92, 0.97),
        "warn": (1, 0.92, 0.61),
        "bad": (1, 0.78, 0.81),
    }

    doc = fitz.open(pdf_path)
    for r in results:
        pagina = r["pdf_pagina"]
        if not pagina or pagina > doc.page_count:
            continue
        page = doc[pagina - 1]
        pw = page.rect.width
        box_w = min(pw - 20, 220)
        rect = fitz.Rect(10, 10, 10 + box_w, 46)
        color = box_colors[categoria_metodo(r["metodo"])]
        page.draw_rect(rect, color=(0, 0, 0), fill=color, width=1)
        page.insert_textbox(
            rect,
            f"CODIGO: {r['codigo']}",
            fontsize=15,
            fontname="helv",
            color=(0, 0, 0),
            align=fitz.TEXT_ALIGN_CENTER,
        )
    doc.save(out_pdf_path)
    doc.close()


def run_matching(xls_path, pdf_path, out_xlsx_path, out_pdf_path, workdir, progress_cb=None):
    pdf_entries = ocr_pdf(pdf_path, workdir, progress_cb=progress_cb)
    excel_rows = load_excel(xls_path)
    results = match(excel_rows, pdf_entries)
    resumen = write_output(results, out_xlsx_path)
    annotate_pdf(pdf_path, results, out_pdf_path)
    return {
        "n_pdf": len(pdf_entries),
        "n_excel": len(excel_rows),
        "resumen": resumen,
    }
